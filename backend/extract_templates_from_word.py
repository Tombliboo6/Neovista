#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从Word文档提取模版数据并生成 templates_v2.json
"""
from docx import Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.text.paragraph import Paragraph
from docx.table import Table
import json
import re
import os
from io import BytesIO
from PIL import Image

def extract_categories_and_templates(doc):
    """提取分类和模版结构"""
    # 手动定义10个大分类（从Word文档中提取）
    categories = {
        '1': '场地分析',
        '2': '环境与物理性能分析',
        '3': '概念与体块推演',
        '4': '功能与程序分析',
        '5': '流线与动线分析',
        '6': '技术与构造分析',
        '7': '感官与用户体验分析',
        '8': '植物分析',
        '9': '其他分析',
        '10': '平面设计'
    }

    subcategories = {}
    templates = []
    tips_map = {}
    prompts_map = {}  # 新增：存储提示词

    current_subcategory = None
    current_template_id = None
    collecting_prompt = False
    current_prompt = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # 子分类 (1.1 城市肌理图)
        subcat_match = re.match(r'^(\d+\.\d+)\s*(.+)$', text)
        if subcat_match and not re.match(r'^\d+\.\d+\.\d+', text):
            subcat_id = subcat_match.group(1)
            subcat_name = subcat_match.group(2)
            subcategories[subcat_id] = subcat_name
            current_subcategory = subcat_id
            continue

        # 模版 (1.1.1 BIG风)
        template_match = re.match(r'^(\d+\.\d+\.\d+)\s*(.*)$', text)
        if template_match:
            # 保存上一个模版的提示词
            if current_template_id and current_prompt:
                prompts_map[current_template_id] = '\n'.join(current_prompt).strip()
                current_prompt = []

            template_id = template_match.group(1)
            template_title = template_match.group(2).strip()

            cat_id = template_id.split('.')[0]
            subcat_id = '.'.join(template_id.split('.')[:2])

            templates.append({
                'id': template_id,
                'title': template_title,
                'category_id': cat_id,
                'category_name': categories.get(cat_id, ''),
                'subcategory_id': subcat_id,
                'subcategory_name': subcategories.get(subcat_id, '')
            })
            current_template_id = template_id
            collecting_prompt = False
            continue

        # Tips
        if current_template_id and (text.lower().startswith('tips') or text.lower().startswith('tip')):
            tips_text = text.replace('Tips：', '').replace('Tips:', '').replace('tips:', '').replace('Tip:', '').strip()
            tips_map[current_template_id] = tips_text
            continue

        # 提示词开始
        if current_template_id and '提示词' in text and text == '提示词':
            collecting_prompt = True
            continue

        # 收集提示词内容
        if collecting_prompt and current_template_id:
            # 遇到下一个模版ID或子分类，停止收集
            if re.match(r'^\d+\.\d+', text):
                collecting_prompt = False
            else:
                current_prompt.append(text)

    # 保存最后一个模版的提示词
    if current_template_id and current_prompt:
        prompts_map[current_template_id] = '\n'.join(current_prompt).strip()

    return categories, subcategories, templates, tips_map, prompts_map

def extract_images(doc, templates, output_dir):
    """提取模版图片"""
    os.makedirs(output_dir, exist_ok=True)
    template_images = {}
    current_id = None
    found_effect = False
    found_prompt = False

    # 特殊模版：允许多个效果图区域
    special_templates = {'9.3.3', '9.3.4', '9.3.5', '9.3.6'}
    debug_ids = special_templates

    for idx, element in enumerate(doc.element.body):
        if isinstance(element, CT_P):
            para = Paragraph(element, doc)
            text = para.text.strip()

            match = re.match(r'^(\d+\.\d+\.\d+)', text)
            if match:
                current_id = match.group(1)
                found_effect = False
                found_prompt = False
                if current_id in debug_ids:
                    print(f"\n[{idx}] 找到{current_id}")
                continue

            if current_id and current_id in debug_ids:
                if '效果图' in text:
                    print(f"[{idx}] 效果图 -> found_effect=True")
                    found_effect = True
                    continue
                if '提示词' in text:
                    print(f"[{idx}] 提示词 -> found_effect=False, found_prompt={not (current_id in special_templates)}")
                    found_effect = False
                    if current_id not in special_templates:
                        found_prompt = True
                    continue
                if found_effect:
                    has_pic = bool(para.runs and any(r._element.xpath('.//pic:pic') for r in para.runs))
                    if has_pic:
                        print(f"[{idx}] 发现图片！")
            elif current_id:
                if '效果图' in text:
                    found_effect = True
                    continue
                if '提示词' in text:
                    found_effect = False
                    if current_id not in special_templates:
                        found_prompt = True
                    continue

            # 普通模版：found_prompt 后停止；特殊模版：继续
            should_extract = current_id and found_effect and (current_id in special_templates or not found_prompt)

            if should_extract:
                for run in para.runs:
                    if run._element.xpath('.//pic:pic'):
                        save_images_from_run(run, current_id, template_images, output_dir)

        elif isinstance(element, CT_Tbl):
            should_extract = current_id and found_effect and (current_id in special_templates or not found_prompt)
            if should_extract:
                table = Table(element, doc)
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            for run in para.runs:
                                if run._element.xpath('.//pic:pic'):
                                    save_images_from_run(run, current_id, template_images, output_dir)

    return template_images

def save_images_from_run(run, template_id, template_images, output_dir):
    """从run中保存图片 - 精确提取"""
    try:
        # 直接从run的XML中提取图片的embed ID
        blip_elements = run._element.xpath('.//a:blip')

        for blip in blip_elements:
            embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            if embed_id and embed_id in run.part.related_parts:
                image_part = run.part.related_parts[embed_id]
                img_data = image_part.blob

                if template_id not in template_images:
                    template_images[template_id] = []

                idx = len(template_images[template_id])
                filename = f"{template_id.replace('.', '_')}_{idx}.png"
                filepath = os.path.join(output_dir, filename)

                with open(filepath, 'wb') as f:
                    f.write(img_data)
                template_images[template_id].append(f"/static/template_images/{filename}")
    except Exception as e:
        print(f"警告: 保存图片失败 {template_id}: {e}")

def detect_multi_step(doc, templates):
    """检测分步骤模版"""
    multi_step_ids = set()
    current_id = None
    
    for para in doc.paragraphs:
        text = para.text.strip()
        match = re.match(r'^(\d+\.\d+\.\d+)', text)
        if match:
            current_id = match.group(1)
        elif current_id and re.search(r'step\s*[.:]?\s*\d+', text, re.IGNORECASE):
            multi_step_ids.add(current_id)
    
    return multi_step_ids

def match_existing_data(word_templates, existing_json):
    """智能匹配现有JSON数据"""
    matched = {}
    
    for wt in word_templates:
        wid = wt['id']
        # 优先ID匹配
        for et in existing_json:
            if et['id'] == wid:
                matched[wid] = et
                break
        
        # 标题相似度匹配
        if wid not in matched:
            for et in existing_json:
                if wt['title'] and et.get('title') and wt['title'] in et['title']:
                    matched[wid] = et
                    break
    
    return matched

def main():
    doc_path = '../提示词模版库10326.docx'
    output_dir = 'static/template_images'
    existing_json_path = 'templates_final.json'
    output_json_path = 'templates_v2.json'
    
    print("📖 读取Word文档...")
    doc = Document(doc_path)

    print("📋 提取分类和模版...")
    categories, subcategories, templates, tips_map, prompts_map = extract_categories_and_templates(doc)
    
    print("🖼️  提取图片...")
    template_images = extract_images(doc, templates, output_dir)
    
    print("🔍 检测分步骤模版...")
    multi_step_ids = detect_multi_step(doc, templates)
    
    print("🔗 匹配现有数据...")
    with open(existing_json_path, 'r', encoding='utf-8') as f:
        existing_data = json.load(f)
    
    matched_data = match_existing_data(templates, existing_data)
    
    print("📦 生成新JSON...")
    # 按ID排序
    templates.sort(key=lambda x: [int(n) for n in x['id'].split('.')])

    result = []
    for t in templates:
        tid = t['id']
        existing = matched_data.get(tid, {})
        
        # 检查图片
        images = template_images.get(tid, [])
        if not images and existing.get('images'):
            images = existing['images']

        # 标题格式：子分类名 - 模版名
        display_title = f"{t['subcategory_name']} - {t['title']}" if t['title'] else t['subcategory_name']

        result.append({
            'id': tid,
            'title': display_title,
            'category_id': t['category_id'],
            'category_name': t['category_name'],
            'subcategory_id': t['subcategory_id'],
            'subcategory_name': t['subcategory_name'],
            'tips': tips_map.get(tid, ''),
            'images': images,
            'is_i2i': existing.get('is_i2i', False),
            'is_multi_step': tid in multi_step_ids,
            'prompt_structure': existing.get('prompt_structure', {}),
            'real_prompt': prompts_map.get(tid, existing.get('real_prompt', '')),
            'display_text': existing.get('display_text', ''),
            'likes': existing.get('likes', 0),
            'uses': existing.get('uses', 0)
        })
    
    with open(output_json_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    
    print(f"\n✅ 完成！")
    print(f"总模版: {len(result)}")
    print(f"有图片: {sum(1 for t in result if t['images'])}")
    print(f"分步骤: {len(multi_step_ids)}")
    print(f"有Tips: {len(tips_map)}")

if __name__ == '__main__':
    main()
